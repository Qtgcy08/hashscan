# 导入必要的库和模块
import os  # 操作系统接口
import sys  # 系统相关参数和函数
import argparse  # 命令行参数解析
import hashlib  # 安全哈希和消息摘要
import threading  # 线程支持
import traceback  # 异常堆栈跟踪
import psutil  # 系统监控
import time  # 时间访问和转换
import signal  # 信号处理
import multiprocessing  # 多进程支持
import json  # JSON编码解码
import yaml  # YAML格式处理
from queue import Queue, Empty  # 队列数据结构
from pathlib import Path  # 面向对象的文件系统路径
from collections import defaultdict  # 带默认值的字典
from concurrent.futures import ThreadPoolExecutor  # 线程池执行器
from rich.console import Console  # 富文本控制台输出
from rich.table import Table  # 富文本表格
from rich.progress import Progress, BarColumn, TextColumn, TimeElapsedColumn, TimeRemainingColumn  # 进度条组件

# 初始化富文本控制台
console = Console()
# 定义队列最大和最小大小常量
MAX_QUEUE_SIZE = 5000
MIN_QUEUE_SIZE = 500
# 定义终止信号常量
TERMINATE_SIGNAL = "TERMINATE"
# 创建线程事件对象
stop_event = threading.Event()  # 停止事件
pause_event = threading.Event()  # 暂停事件


class DynamicQueue(Queue):
    """动态调整大小的队列实现"""
    def __init__(self, maxsize=0):
        super().__init__(maxsize)
        self._dynamic_maxsize = maxsize  # 动态最大大小
        self._size_lock = threading.Lock()  # 大小调整锁

    def set_maxsize(self, new_maxsize):
        """设置队列的新大小"""
        with self._size_lock:
            self._dynamic_maxsize = new_maxsize
            self.maxsize = new_maxsize
            with self.not_full:
                self.not_full.notify_all()  # 通知所有等待的线程

    @property
    def safe_qsize(self):
        """线程安全获取队列大小"""
        with self._size_lock:
            return self.qsize()


class DynamicController:
    """动态队列控制器，根据系统负载调整队列大小"""
    def __init__(self, file_queue):
        self.file_queue = file_queue  # 控制的队列对象
        self.adjust_interval = 3  # 调整间隔(秒)
        self.cpu_threshold = 75  # CPU使用率阈值(%)
        self.mem_threshold = 70  # 内存使用率阈值(%)
        self.queue_util_high = 0.7  # 队列利用率高阈值
        self.queue_util_low = 0.3  # 队列利用率低阈值
        self.last_cpu = 0  # 上次记录的CPU使用率
        self.last_mem = 0  # 上次记录的内存使用率

    def get_system_load(self):
        """获取当前系统负载(CPU和内存)"""
        try:
            return psutil.cpu_percent(interval=0.1), psutil.virtual_memory().percent
        except Exception:
            return 0, 0  # 出错时返回0值

    def calculate_size(self):
        """计算新的队列大小"""
        current_size = self.file_queue._dynamic_maxsize
        qsize = self.file_queue.safe_qsize
        if current_size <= 0:
            return MAX_QUEUE_SIZE  # 初始大小设为最大值
        
        # 获取当前系统负载
        self.last_cpu, self.last_mem = self.get_system_load()
        utilization = qsize / current_size if current_size else 0  # 计算队列利用率
        
        # 根据系统负载和队列利用率动态调整大小
        if self.last_cpu > self.cpu_threshold or self.last_mem > self.mem_threshold:
            return max(MIN_QUEUE_SIZE, int(current_size * 0.7))  # 负载高时缩小队列
        return (
            min(MAX_QUEUE_SIZE, int(current_size * 1.3)) if utilization > self.queue_util_high
            else max(MIN_QUEUE_SIZE, int(current_size * 0.8)) if utilization < self.queue_util_low
            else current_size  # 根据利用率调整
        )

    def adjust_loop(self):
        """队列大小调整循环"""
        while not stop_event.is_set():
            time.sleep(self.adjust_interval)
            new_size = self.calculate_size()
            current_size = self.file_queue._dynamic_maxsize
            if new_size != current_size:
                console.print(
                    f"队列调整: {current_size} → {new_size} (CPU: {self.last_cpu}% MEM: {self.last_mem}%)",
                    style="bold yellow"
                )
                self.file_queue.set_maxsize(new_size)


def signal_handler(sig, frame):
    """信号处理函数，处理Ctrl+C中断"""
    console.print("\n操作被用户中断", style="bold red")
    stop_event.set()  # 设置停止事件
    sys.exit(0)


# 注册信号处理
signal.signal(signal.SIGINT, signal_handler)


class HashScanner:
    """文件哈希扫描器主类"""
    def __init__(self, folder, algorithms, recursive, threads):
        self.folder = Path(folder).resolve()  # 扫描目录
        self.algorithms = [a.lower() for a in algorithms]  # 哈希算法列表
        self.recursive = recursive  # 是否递归扫描
        self.threads = threads  # 线程数
        self.file_queue = DynamicQueue(maxsize=MAX_QUEUE_SIZE)  # 文件队列
        self.controller = DynamicController(self.file_queue)  # 队列控制器
        self.monitor_thread = threading.Thread(target=self.controller.adjust_loop, daemon=True)  # 监控线程
        self.hash_results = defaultdict(dict)  # 哈希结果存储
        self.total_files = 0  # 总文件数
        self.processed_files = 0  # 已处理文件数
        self.error_count = 0  # 错误计数
        self.total_lock = threading.Lock()  # 总文件数锁
        self.progress_lock = threading.Lock()  # 进度锁
        self.thread_status = {i+1: None for i in range(threads)}  # 线程状态字典
        self.status_lock = threading.Lock()  # 状态锁

    def _validate_folder(self):
        """验证目标文件夹有效性"""
        if not self.folder.exists():
            raise FileNotFoundError(f"目标文件夹不存在: {self.folder}")
        if not self.folder.is_dir():
            raise NotADirectoryError(f"路径不是目录: {self.folder}")

    def _file_producer(self):
        """文件生产者线程，遍历文件系统"""
        try:
            scan_method = os.walk if self.recursive else os.scandir  # 选择扫描方法
            for entry in scan_method(self.folder):
                if stop_event.is_set():  # 检查停止信号
                    break
                if self.recursive:
                    root, _, files = entry
                    for f in files:
                        self._add_to_queue(Path(root) / f)  # 递归模式下添加文件
                else:
                    if entry.is_file():
                        self._add_to_queue(Path(entry.path))  # 非递归模式下添加文件
        except Exception as e:
            self._log_error(f"文件遍历错误: {str(e)}")
            stop_event.set()  # 出错时设置停止事件
        finally:
            for _ in range(self.threads):
                self.file_queue.put(TERMINATE_SIGNAL)  # 发送终止信号

    def _add_to_queue(self, path):
        """添加文件到队列"""
        try:
            if not path.exists():
                raise FileNotFoundError("文件不存在")
            if path.is_symlink():
                console.print(f"[dim]跳过符号链接: {path}[/dim]")  # 跳过符号链接
                return
            if not os.access(path, os.R_OK):
                raise PermissionError("无读取权限")  # 检查读取权限
            with self.total_lock:
                self.total_files += 1  # 原子操作增加总文件数
            self.file_queue.put(str(path), timeout=5)  # 添加文件路径到队列
        except Exception as e:
            self._log_error(f"文件添加失败 [{type(e).__name__}]: {path} 原因: {str(e)}")
            with self.progress_lock:
                self.error_count += 1  # 增加错误计数

    def _hash_consumer(self, thread_id):
        """哈希消费者线程，计算文件哈希"""
        try:
            while not stop_event.is_set():
                if pause_event.is_set():  # 检查暂停状态
                    time.sleep(0.5)
                    continue
                try:
                    path_str = self.file_queue.get(timeout=0.5)  # 从队列获取文件路径
                    if path_str == TERMINATE_SIGNAL:  # 检查终止信号
                        break
                    # 处理Windows长路径前缀
                    path = Path(path_str[4:] if os.name == 'nt' and path_str.startswith('\\\\?\\') else path_str)
                    with self.status_lock:
                        self.thread_status[thread_id] = str(path)  # 更新线程状态
                    try:
                        with open(path, 'rb') as f:
                            # 初始化哈希器
                            hashers = {algo: hashlib.new(algo) for algo in self.algorithms}
                            # 分块读取文件计算哈希
                            for chunk in iter(lambda: f.read(131072), b''):
                                mv = memoryview(chunk)  # 使用内存视图提高性能
                                for h in hashers.values():
                                    h.update(mv)
                            # 存储哈希结果
                            self.hash_results[path] = {algo: h.hexdigest() for algo, h in hashers.items()}
                    except Exception as e:
                        self.hash_results[path] = f"Error: {str(e)}"  # 存储错误信息
                        with self.progress_lock:
                            self.error_count += 1  # 增加错误计数
                    with self.progress_lock:
                        self.processed_files += 1  # 增加已处理计数
                except Empty:
                    continue  # 队列为空时继续
                finally:
                    with self.status_lock:
                        self.thread_status[thread_id] = None  # 重置线程状态
        except Exception as e:
            self._log_error(f"工作线程异常: {str(e)}")

    def _log_error(self, message):
        """错误日志记录"""
        console.print(f"[bold red]错误: {message}[/bold red]")

    def run(self):
        """执行扫描过程"""
        self._validate_folder()  # 验证文件夹
        self.monitor_thread.start()  # 启动监控线程

        def count_total_files():
            """计算总文件数"""
            count = 0
            try:
                scan_method = os.walk if self.recursive else os.scandir
                for entry in scan_method(self.folder):
                    if self.recursive:
                        root, _, files = entry
                        count += len(files)  # 递归模式下计数
                    else:
                        if entry.is_file():
                            count += 1  # 非递归模式下计数
            except Exception as e:
                console.print(f"[bold red]文件计数错误: {str(e)}[/bold red]")
            with self.total_lock:
                self.total_files = count  # 更新总文件数

        # 启动文件计数线程
        counter_thread = threading.Thread(target=count_total_files, daemon=True)
        counter_thread.start()
        counter_thread.join(timeout=10)  # 等待计数完成或超时

        progress_lock = threading.Lock()
        # 使用线程池执行任务
        with ThreadPoolExecutor(max_workers=self.threads + 2) as executor:
            producer = executor.submit(self._file_producer)  # 提交生产者任务
            consumers = [executor.submit(self._hash_consumer, i+1) for i in range(self.threads)]  # 提交消费者任务

            try:
                # 创建进度条
                with Progress(
                    TextColumn("[progress.description]{task.description}"),
                    BarColumn(complete_style="blue", finished_style="green"),
                    TextColumn("[progress.percentage]{task.percentage:>3.0f}%"),
                    TextColumn("• 已处理: {task.completed}/{task.total}"),
                    TextColumn("• 错误: [red]{task.fields[errors]}"),
                    TimeElapsedColumn(),
                    TextColumn("•"),
                    TimeRemainingColumn(),
                    console=console,
                    transient=False,
                    refresh_per_second=10
                ) as progress:
                    total = max(self.total_files, 1)  # 确保总数不为零
                    task = progress.add_task(
                        "初始化扫描...",
                        total=total,
                        errors=0
                    )
                    last_processed = 0
                    last_time = time.time()

                    # 进度更新循环
                    while not all(f.done() for f in consumers) and not stop_event.is_set():
                        time.sleep(0.2)
                        current_time = time.time()
                        time_diff = current_time - last_time
                        if time_diff < 0.001:
                            last_time = current_time
                            continue

                        # 更新进度条
                        with progress_lock:
                            progress.update(
                                task,
                                completed=self.processed_files,
                                total=total,
                                description=f"处理进度 | CPU: {self.controller.last_cpu}% | MEM: {self.controller.last_mem}%",
                                errors=self.error_count
                            )
                        last_processed = self.processed_files
                        last_time = current_time

                        # 在终端中显示线程状态表
                        if sys.stdout.isatty():
                            with self.status_lock:
                                table = Table(title="线程状态", show_header=True)
                                table.add_column("线程ID", style="cyan")
                                table.add_column("当前文件", overflow="fold", max_width=40)
                                for tid, file_path in self.thread_status.items():
                                    display_path = Path(file_path).name[:35] if file_path else "空闲"
                                    table.add_row(f"{tid}", display_path)
                                console.print(f"总文件数: {self.total_files} | 已处理: {self.processed_files} | 错误: {self.error_count}")
                                console.print(table)

                    # 最终更新进度条
                    with progress_lock:
                        progress.update(task, completed=self.processed_files)

            except Exception as e:
                console.print(f"[bold red]进度条渲染异常: {str(e)}[/bold red]")
                traceback.print_exc()

        return self.hash_results  # 返回哈希结果

    def _handle_input(self):
        """处理用户输入命令"""
        while not stop_event.is_set():
            try:
                cmd = input().strip().lower()  # 获取用户输入
                if cmd == 'p':
                    pause_event.set()  # 暂停扫描
                    console.print("[bold yellow]扫描已暂停[/bold yellow]")
                elif cmd == 'c':
                    pause_event.clear()  # 继续扫描
                    console.print("[bold green]扫描已继续[/bold green]")
                elif cmd == 'q':
                    stop_event.set()  # 停止扫描
                    console.print("[bold red]扫描已终止[/bold red]")
            except EOFError:
                pass


def format_results(results, args, scanner):
    """格式化输出结果"""
    if args.same:  # 仅显示重复文件模式
        hash_groups = defaultdict(list)
        for path, hashes in results.items():
            if isinstance(hashes, dict):
                key = tuple(sorted(hashes.items()))  # 创建哈希键
                hash_groups[key].append(str(path))  # 分组相同哈希文件
        return _format_duplicates(hash_groups, args.format, scanner)
    return _format_normal(results, args.format, args.algorithms, scanner)  # 普通模式


def _format_normal(results, fmt, algorithms, scanner):
    """格式化普通结果"""
    if fmt == 'xlsx':  # Excel格式输出
        from openpyxl import Workbook
        from openpyxl.utils import get_column_letter
        from openpyxl.styles import Font, Alignment
        import tempfile

        wb = Workbook()
        ws = wb.active
        ws.title = "文件哈希值"

        # 写入元数据
        meta_data = [
            ["扫描目录:", str(scanner.folder)],
            ["扫描模式:", "递归" if scanner.recursive else "非递归"],
            ["哈希算法:", ", ".join(algorithms)],
            ["总文件数:", scanner.total_files],
            ["成功处理:", scanner.processed_files - scanner.error_count],
            ["错误文件:", scanner.error_count],
            ["扫描时间:", time.strftime('%Y-%m-%d %H:%M:%S')],
            []  # 空行
        ]

        for row in meta_data:
            ws.append(row)

        # 写入表头
        headers = ["文件路径"] + [a.upper() for a in algorithms]
        ws.append(headers)

        # 设置表头样式
        header_font = Font(bold=True)
        for cell in ws[ws.max_row]:
            cell.font = header_font

        # 写入数据
        for path, hashes in results.items():
            row = [str(path)]
            if isinstance(hashes, dict):
                row.extend([hashes[a] for a in algorithms])  # 添加哈希值
            else:
                row.append(hashes)  # 添加错误信息
            ws.append(row)

        # 调整列宽和样式
        for col in range(1, len(headers) + 1):
            col_letter = get_column_letter(col)
            if col == 1:  # 文件路径列
                ws.column_dimensions[col_letter].width = 50
            else:
                ws.column_dimensions[col_letter].width = 20

            # 设置自动换行
            for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=col, max_col=col):
                for cell in row:
                    cell.alignment = Alignment(wrap_text=True, vertical='top')

        # 冻结首行
        ws.freeze_panes = 'A2'

        # 保存到临时文件
        temp_path = tempfile.mktemp(suffix='.xlsx')
        wb.save(temp_path)

        # 读取二进制内容返回
        with open(temp_path, 'rb') as f:
            content = f.read()
        
        # 删除临时文件
        try:
            os.unlink(temp_path)
        except:
            pass

        return content

    elif fmt == 'json':  # JSON格式输出
        output = {
            "#meta": {
                "directory": str(scanner.folder),
                "mode": "recursive" if scanner.recursive else "non-recursive",
                "algorithms": algorithms,
                "total_files": scanner.total_files,
                "processed": scanner.processed_files,
                "errors": scanner.error_count,
                "timestamp": time.strftime('%Y-%m-%d %H:%M:%S')
            },
            "results": {str(k): v for k, v in results.items()}
        }
        return json.dumps(output, indent=2, ensure_ascii=False)
    
    elif fmt == 'yaml':  # YAML格式输出
        output = {
            "meta": {
                "directory": str(scanner.folder),
                "mode": "recursive" if scanner.recursive else "non-recursive",
                "algorithms": algorithms,
                "total_files": scanner.total_files,
                "processed": scanner.processed_files,
                "errors": scanner.error_count,
                "timestamp": time.strftime('%Y-%m-%d %H:%M:%S')
            },
            "results": {str(k): v for k, v in results.items()}
        }
        return yaml.dump(output, default_flow_style=False, sort_keys=False, allow_unicode=True)
    
    elif fmt == 'markdown':  # Markdown格式输出
        md = "# 文件哈希扫描报告\n\n"
        md += f"- **扫描目录**: {scanner.folder}\n"
        md += f"- **扫描模式**: {'递归' if scanner.recursive else '非递归'}\n"
        md += f"- **哈希算法**: {', '.join(algorithms)}\n"
        md += f"- **总文件数**: {scanner.total_files}\n"
        md += f"- **成功处理**: {scanner.processed_files - scanner.error_count}\n"
        md += f"- **错误文件**: {scanner.error_count}\n\n"
        md += "---\n\n"
        md += "| 文件路径 | " + " | ".join(a.upper() for a in algorithms) + " |\n"
        md += "|" + "|".join(["---"] * (len(algorithms)+1)) + "|\n"
        for path, hashes in results.items():
            if isinstance(hashes, dict):
                md += f"| {str(path)} | " + " | ".join(hashes[a] for a in algorithms) + " |\n"
            else:
                md += f"| {str(path)} | {hashes} |\n"
        md += f"\n> 扫描完成于: {time.strftime('%Y-%m-%d %H:%M:%S')}"
        return md
    
    else:  # 默认文本格式输出
        output = []
        output.append("文件哈希扫描报告")
        output.append("=" * 50)
        output.append(f"扫描目录: {scanner.folder}")
        output.append(f"扫描模式: {'递归' if scanner.recursive else '非递归'}")
        output.append(f"哈希算法: {', '.join(algorithms)}")
        output.append(f"总文件数: {scanner.total_files}")
        output.append(f"成功处理: {scanner.processed_files - scanner.error_count}")
        output.append(f"错误文件: {scanner.error_count}")
        output.append("=" * 50 + "\n")

        for path, hashes in results.items():
            output.append(f"文件: {path}")
            if isinstance(hashes, dict):
                for algo, digest in hashes.items():
                    output.append(f"  {algo.upper()}: {digest}")
            else:
                output.append(f"  错误: {hashes}")
            output.append("")  # 空行分隔
        
        output.append("=" * 50)
        output.append(f"扫描完成于: {time.strftime('%Y-%m-%d %H:%M:%S')}")
        return "\n".join(output)


def _format_duplicates(hash_groups, fmt, scanner):
    """格式化重复文件结果"""
    groups = [{"hashes": dict(k), "files": v} for k, v in hash_groups.items() if len(v) > 1]  # 筛选重复文件组
    
    if fmt == 'xlsx':  # Excel格式输出
        from openpyxl import Workbook
        from openpyxl.utils import get_column_letter
        from openpyxl.styles import Font, Alignment
        import tempfile

        if not groups:
            return "未发现重复文件"

        wb = Workbook()
        ws = wb.active
        ws.title = "重复文件"

        # 写入元数据
        meta_data = [
            ["扫描目录:", str(scanner.folder)],
            ["扫描模式:", "递归" if scanner.recursive else "非递归"],
            ["哈希算法:", ", ".join(scanner.algorithms)],
            ["总文件数:", scanner.total_files],
            ["成功处理:", scanner.processed_files - scanner.error_count],
            ["错误文件:", scanner.error_count],
            ["扫描时间:", time.strftime('%Y-%m-%d %H:%M:%S')],
            ["重复文件组数:", len(groups)],
            []  # 空行
        ]

        for row in meta_data:
            ws.append(row)

        # 写入表头
        headers = ["哈希类型", "哈希值", "重复文件数", "文件列表"]
        ws.append(headers)

        # 设置表头样式
        header_font = Font(bold=True)
        for cell in ws[ws.max_row]:
            cell.font = header_font

        # 写入数据
        for group in groups:
            for algo, digest in group["hashes"].items():
                files = "\n".join(f"- {f}" for f in group["files"])  # 格式化文件列表
                ws.append([algo.upper(), digest, len(group["files"]), files])

        # 调整列宽和样式
        col_widths = {"A": 15, "B": 40, "C": 15, "D": 60}
        for col, width in col_widths.items():
            ws.column_dimensions[col].width = width

        # 设置自动换行
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=4):
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical='top')

        # 冻结首行
        ws.freeze_panes = 'A2'

        # 保存到临时文件
        temp_path = tempfile.mktemp(suffix='.xlsx')
        wb.save(temp_path)

        # 读取二进制内容返回
        with open(temp_path, 'rb') as f:
            content = f.read()
        
        # 删除临时文件
        try:
            os.unlink(temp_path)
        except:
            pass

        return content
    
    elif fmt == 'json':  # JSON格式输出
        return json.dumps(groups, indent=2, ensure_ascii=False)
    
    elif fmt == 'yaml':  # YAML格式输出
        return yaml.dump(groups, default_flow_style=False, sort_keys=False, allow_unicode=True)
    
    elif fmt == 'markdown':  # Markdown格式输出
        if not groups:
            return "## 重复文件报告\n\n未发现重复文件"
        
        md = "## 重复文件报告\n\n"
        for group in groups:
            md += "### 重复文件组\n"
            for algo, digest in group["hashes"].items():
                md += f"- **{algo.upper()}**: `{digest}`\n"
            md += f"- **重复文件数**: {len(group['files'])}\n"
            md += "\n**包含文件**:\n"
            for f in group["files"]:
                md += f"  - {f}\n"
            md += "\n"
        return md
    
    else:  # 默认文本格式输出
        if not groups:
            return "重复文件报告\n" + "="*50 + "\n未发现重复文件"
        
        output = ["重复文件报告", "=" * 50]
        for group in groups:
            output.append(f"重复文件组 (共{len(group['files'])}个):")
            for algo, digest in group["hashes"].items():
                output.append(f"  {algo.upper()}: {digest}")
            output.append("  包含文件:")
            for f in group["files"]:
                output.append(f"    - {f}")
            output.append("")  # 空行分隔
        output.append("=" * 50)
        return "\n".join(output)


def main():
    """主函数"""
    # 检查是否通过拖放文件方式运行
    if len(sys.argv) > 1 and not sys.argv[1].startswith('-'):
        # 如果是拖放文件方式运行，自动添加参数
        folder_path = sys.argv[1]
        sys.argv = [sys.argv[0], '-f', folder_path, '-a', 'md5', 'sha1', 'sha256', 'sha512', '-r']
    
    # 创建自定义格式的帮助文本
    class CustomHelpFormatter(argparse.HelpFormatter):
        """自定义帮助格式"""
        def __init__(self, prog):
            super().__init__(prog, max_help_position=40, width=100)
        
        def _format_action_invocation(self, action):
            """格式化参数调用显示"""
            if not action.option_strings:
                return self._metavar_formatter(action, action.dest)(1)[0]
            else:
                parts = []
                # 显示所有选项，用逗号分隔
                if len(action.option_strings) > 1:
                    parts.append(', '.join(action.option_strings))
                else:
                    parts.append(action.option_strings[0])
                # 如果有参数，添加参数
                if action.nargs != 0:
                    default = self._get_default_metavar_for_optional(action)
                    args_string = self._format_args(action, default)
                    parts.append(' %s' % args_string)
                return ''.join(parts)
    
    # 创建参数解析器
    parser = argparse.ArgumentParser(
        description="智能文件哈希扫描工具",
        formatter_class=CustomHelpFormatter,
        epilog="使用示例:\n"
               "  直接运行: python hash_scanner.py -f /path/to/folder\n"
               "  拖放文件夹: 将文件夹拖放到脚本文件上\n"
               "  生成Excel报告: python hash_scanner.py -f /path/to/folder --format xlsx -o report.xlsx\n"
               "  查找重复文件: python hash_scanner.py -f /path/to/folder -s\n\n"
               "支持的哈希算法: " + ", ".join(sorted(hashlib.algorithms_available)))
    
    # 参数分组
    required_args = parser.add_argument_group('必需参数')
    optional_args = parser.add_argument_group('可选参数')
    output_args = parser.add_argument_group('输出选项')
    
    # 必需参数
    required_args.add_argument(
        '-f', '--folder',
        default='.',
        help="要扫描的目标文件夹路径 (默认: 当前目录)"
    )
    
    # 可选参数
    optional_args.add_argument(
        '-a', '--algorithms',
        nargs='+',
        default=['md5'],
        choices=hashlib.algorithms_available,
        metavar='ALGO',
        help="要使用的哈希算法 (默认: md5)"
    )
    optional_args.add_argument(
        '-r', '--recursive',
        action='store_true',
        help="递归扫描子目录"
    )
    optional_args.add_argument(
        '-t', '--threads',
        type=int,
        default=multiprocessing.cpu_count(),
        help="使用的线程数 (默认: CPU核心数)"
    )
    optional_args.add_argument(
        '-s', '--same',
        action='store_true',
        help="仅显示重复文件"
    )
    
    # 输出选项
    output_args.add_argument(
        '-o', '--output',
        help="将结果输出到指定文件"
    )
    output_args.add_argument(
        '--format',
        choices=['text', 'json', 'yaml', 'markdown', 'xlsx'],
        default='text',
        help="输出格式 (默认: text)"
    )
    
    # 如果没有提供参数，显示帮助
    if len(sys.argv) == 1:
        parser.print_help()
        sys.exit(0)
    
    args = parser.parse_args()  # 解析参数


    try:
        # 初始化扫描器
        scanner = HashScanner(args.folder, args.algorithms, args.recursive, args.threads)
        console.print(f"初始队列: {MAX_QUEUE_SIZE}", style="bold blue")
        console.print(f"工作线程: {args.threads}", style="bold blue")
        
        # 启动输入处理线程
        input_handler = threading.Thread(target=scanner._handle_input, daemon=True)
        input_handler.start()
        
        # 执行扫描并获取结果
        results = scanner.run()
        formatted = format_results(results, args, scanner)

        # 处理输出
        if args.output:
            if args.format == 'xlsx':
                with open(args.output, 'wb') as f:  # 二进制模式写入Excel文件
                    f.write(formatted)
            else:
                with open(args.output, 'w', encoding='utf-8') as f:
                    f.write(formatted)
            console.print(f"\n[bold green]结果已保存到: {args.output}[/bold green]")
        else:
            if args.format == 'xlsx':
                console.print("[bold red]错误: xlsx格式需要指定输出文件路径(-o/--output)[/bold red]")
                sys.exit(1)
            console.print("\n" + formatted)  # 直接输出到控制台

    except Exception as e:
        console.print(f"[bold red]致命错误: {traceback.format_exc()}[/bold red]")
        sys.exit(1)

    finally:
        stop_event.set()  # 确保停止事件被设置


if __name__ == '__main__':
    main()  # 程序入口